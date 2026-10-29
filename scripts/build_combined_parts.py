"""Build small part files and a master file for a unified Word report."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from build_course_docs import (  # noqa: E402
    DOCUMENTS,
    OUTPUT_DIR,
    SOURCE_DIR,
    add_cover,
    add_header_footer,
    add_inline,
    add_markdown,
    add_toc,
    style_document,
)
from docx import Document  # noqa: E402

BUILD_DIR = OUTPUT_DIR / "_build"


def main():
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    master = Document()
    style_document(master)
    add_header_footer(master, "课程大作业完整报告")
    add_cover(master, "课程大作业完整报告", "V1.0")
    add_toc(master)
    master.save(BUILD_DIR / "00-master.docx")

    for index, (filename, title, _) in enumerate(DOCUMENTS, start=1):
        path = SOURCE_DIR / filename
        markdown = path.read_text(encoding="utf-8")
        doc = Document()
        style_document(doc)
        add_header_footer(doc, f"第{index}部分 {title}")
        p = doc.add_paragraph(style="Heading 1")
        add_inline(p, f"第{index}部分  {title}", size=18, bold=True)
        add_markdown(doc, markdown, path)
        output = BUILD_DIR / f"{index:02d}-{title.replace(' ', '-')}.docx"
        doc.save(output)
        print(f"built {output.name}")


if __name__ == "__main__":
    main()
