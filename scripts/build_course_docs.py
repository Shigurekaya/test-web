"""Build polished DOCX deliverables from the course Markdown sources."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.image.image import Image as DocxImage
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SOURCE_DIR = ROOT / "docs" / "course"
OUTPUT_DIR = ROOT / "deliverables"
PROJECT_NAME = "智慧企业知识财富库系统"
COURSE_NAME = "软件工程课程大作业"
DOC_DATE = "2026-10-31"
AUTHOR = "项目负责人（本人）"

DOCUMENTS = [
    ("01-项目开发计划.md", "项目开发计划", "V1.0"),
    ("02-软件需求规格说明书.md", "软件需求规格说明书", "V1.0"),
    ("03-软件设计规格说明书.md", "软件设计规格说明书", "V1.0"),
    ("04-项目测试计划.md", "项目测试计划", "V1.0"),
    ("05-测试用例说明书.md", "测试用例说明书", "V1.0"),
    ("06-项目总结.md", "项目总结", "V1.0"),
    ("07-个人项目小结与考核材料.md", "个人项目小结与考核材料", "V1.0"),
    ("08-REST-API接口设计说明书.md", "REST API 接口设计说明书", "V1.0"),
]

BODY_FONT = "宋体"
HEADING_FONT = "微软雅黑"
MONO_FONT = "Consolas"
BODY_SIZE = 10.5
TABLE_SIZE = 8.5
ACCENT = "1D9B83"
NAVY = "101B33"
LIGHT = "F3F6FA"


def set_run_font(
    run, name: str, size: float | None = None, *, bold: bool | None = None, color: str | None = None
):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_inches: float):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction: str, placeholder: str = ""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)
    if placeholder:
        text = OxmlElement("w:t")
        text.text = placeholder
        run._r.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def add_hyperlink(paragraph, text: str, url: str):
    rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "3976D9")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(color)
    run_props.append(underline)
    run.append(run_props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")


def add_inline(
    paragraph, text: str, *, size: float = BODY_SIZE, color: str | None = None, bold: bool = False
):
    text = text.replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, BODY_FONT, size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, MONO_FONT, max(8.5, size - 1), bold=bold, color=color)
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if not match:
                run = paragraph.add_run(part)
                set_run_font(run, BODY_FONT, size, bold=bold, color=color)
                continue
            label, url = match.groups()
            if url.startswith(("http://", "https://")):
                add_hyperlink(paragraph, label, url)
            else:
                run = paragraph.add_run(label)
                set_run_font(run, BODY_FONT, size, bold=bold, color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, BODY_FONT, size, bold=bold, color=color)


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = True
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.widow_control = True
    for name, size, color in (
        ("Title", 30, NAVY),
        ("Heading 1", 18, NAVY),
        ("Heading 2", 14, NAVY),
        ("Heading 3", 12, "1A6E61"),
        ("Heading 4", 11, "33445B"),
    ):
        style = styles[name]
        style.font.name = HEADING_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    caption = styles["Caption"]
    caption.font.name = BODY_FONT
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string("6E7B90")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)


def add_header_footer(doc: Document, title: str):
    for section in doc.sections:
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_inline(hp, f"{PROJECT_NAME}  |  {title}", size=8.5, color="8793A6")
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(fp, "第 ", size=8.5, color="8793A6")
        add_field(fp, "PAGE", "1")
        add_inline(fp, " 页 / 共 ", size=8.5, color="8793A6")
        add_field(fp, "NUMPAGES", "1")
        add_inline(fp, " 页", size=8.5, color="8793A6")


def add_cover(doc: Document, title: str, version: str = "V1.0"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COURSE_NAME)
    set_run_font(run, HEADING_FONT, 13, bold=True, color=ACCENT)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(46)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(PROJECT_NAME)
    set_run_font(run, HEADING_FONT, 22, bold=True, color="5E6B80")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, HEADING_FONT, 30, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"文档版本 {version}")
    set_run_font(run, BODY_FONT, 11, color="7A879A")
    info = doc.add_table(rows=5, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    info_data = [
        ("课程名称", COURSE_NAME),
        ("项目版本", "V2.0"),
        ("项目组织形式", "单人独立项目组"),
        ("编制人", AUTHOR),
        ("编制日期", DOC_DATE),
    ]
    for index, (key, value) in enumerate(info_data):
        cells = info.rows[index].cells
        set_cell_width(cells[0], 1.8)
        set_cell_width(cells[1], 3.8)
        cells[0].text = ""
        cells[1].text = ""
        add_inline(cells[0].paragraphs[0], key, size=10.5, bold=True, color="6E7B90")
        add_inline(cells[1].paragraphs[0], value, size=10.5)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cells[0], LIGHT)
    info.style = "Table Grid"
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(68)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "课程设计项目 · 演示数据均为虚构", size=9, color="95A0B2")
    doc.add_page_break()


def add_toc(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run("目录")
    set_run_font(run, HEADING_FONT, 20, bold=True, color=NAVY)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "目录将在更新域后生成")


def trim_front_matter(markdown: str) -> str:
    """Drop per-document covers and hand-written TOCs before merging.

    Source documents use either ``# 1`` or ``## 1`` for the first numbered
    section.  The first numbered heading after the front matter is the stable
    body boundary, so both styles (and a leading section 0) are recognised.
    Fenced code is ignored to avoid treating shell comments as headings.
    """
    lines = markdown.splitlines()
    pattern = re.compile(r"^#{1,2}\s+(?:第?\s*\d+(?:[\.\s、]|$)|[一二三四五六七八九十]+[、\.])")
    in_code = False
    for index, line in enumerate(lines):
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            continue
        if not in_code and index > 8 and pattern.match(stripped):
            return "\n".join(lines[index:])
    return markdown


def find_image(src: str, md_path: Path) -> Path | None:
    src = src.strip().strip("<>")
    if src.startswith(("http://", "https://", "data:")):
        return None
    candidate = Path(src)
    if candidate.is_absolute() and candidate.exists():
        return candidate
    for base in (md_path.parent, ROOT):
        resolved = (base / src).resolve()
        if resolved.exists():
            return resolved
    return None


def add_image(doc: Document, md_path: Path, alt: str, src: str):
    image_path = find_image(src, md_path)
    if image_path is None:
        p = doc.add_paragraph()
        add_inline(p, f"[图片未找到：{alt or src}]", size=9, color="B24440")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    max_width = Inches(6.15)
    max_height = Inches(7.2)
    try:
        source = DocxImage.from_file(str(image_path))
        aspect = source.height / source.width
        if max_width * aspect > max_height:
            height = max_height
            width = int(max_height / aspect)
        else:
            width = max_width
            height = int(max_width * aspect)
        run.add_picture(str(image_path), width=width, height=height)
    except Exception:
        run.add_picture(str(image_path), width=Inches(5.4))
    if alt:
        caption = doc.add_paragraph(style="Caption")
        add_inline(caption, alt, size=8.5, color="6E7B90")


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        line = lines[index].strip()
        if re.match(r"^\|?\s*:?-{3,}", line):
            index += 1
            continue
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        rows.append(cells)
        index += 1
    return rows, index


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row_data in enumerate(rows):
        for col_index in range(cols):
            cell = table.cell(row_index, col_index)
            text = row_data[col_index] if col_index < len(row_data) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, text, size=TABLE_SIZE, bold=(row_index == 0))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shade_cell(cell, "E6F6F2")
    repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(doc: Document, lines: list[str]):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, "F6F8FB")
    cell.text = ""
    for index, line in enumerate(lines):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line if line else " ")
        set_run_font(run, MONO_FONT, 8.5, color="344156")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_markdown(
    doc: Document,
    markdown: str,
    md_path: Path,
    *,
    page_break_first: bool = False,
    heading_offset: int = 0,
    normalize_chapter_heading: bool = False,
):
    first_paragraph_index = len(doc.paragraphs)
    lines = trim_front_matter(markdown).splitlines()
    if normalize_chapter_heading:
        chapter_pattern = re.compile(
            r"^#{1,2}\s+(?:第?\s*\d+(?:[\.\s、]|$)|[一二三四五六七八九十]+[、\.])"
        )
        for line in lines:
            stripped = line.strip()
            if chapter_pattern.match(stripped):
                source_level = len(stripped) - len(stripped.lstrip("#"))
                heading_offset = 2 - source_level
                break
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("<!--") or stripped in {
            '<div align="center">',
            "</div>",
            "<div align='center'>",
        }:
            index += 1
            continue
        if stripped in {"---", "***"}:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            index += 1
            continue
        image_match = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", stripped)
        if image_match:
            add_image(doc, md_path, image_match.group(1), image_match.group(2))
            index += 1
            continue
        if stripped.startswith("|"):
            rows, next_index = parse_table(lines, index)
            add_table(doc, rows)
            index = next_index
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)) + heading_offset, 4)
            title = re.sub(r"\s+#+\s*$", "", heading.group(2)).strip()
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, title, size={1: 18, 2: 14, 3: 12, 4: 11}[level], bold=True)
            index += 1
            continue
        if stripped.startswith(">"):
            quote = stripped.lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(5)
            add_inline(p, quote, size=9.5, color="6E7B90")
            index += 1
            continue
        bullet = re.match(r"^[-*+]\s+(.+)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, bullet.group(1))
            index += 1
            continue
        numbered = re.match(r"^\d+[\.\)]\s+(.+)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, numbered.group(1))
            index += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, stripped)
        index += 1
    if in_code and code_lines:
        add_code_block(doc, code_lines)
    if page_break_first and len(doc.paragraphs) > first_paragraph_index:
        doc.paragraphs[first_paragraph_index].paragraph_format.page_break_before = True


def build_document(title: str, version: str, markdown: str, md_path: Path) -> Document:
    doc = Document()
    style_document(doc)
    add_header_footer(doc, title)
    add_cover(doc, title, version)
    add_toc(doc)
    add_markdown(doc, markdown, md_path, page_break_first=True)
    props = doc.core_properties
    props.title = f"{PROJECT_NAME} - {title}"
    props.subject = COURSE_NAME
    props.author = AUTHOR
    props.keywords = "软件工程, 课程大作业, Vue, Flask, RAG, 知识图谱"
    props.comments = "V2.0 课程作业交付文档"
    return doc


def build_all():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    missing = [name for name, _, _ in DOCUMENTS if not (SOURCE_DIR / name).exists()]
    if missing:
        raise FileNotFoundError(f"缺少 Markdown 文档：{', '.join(missing)}")

    for filename, title, version in DOCUMENTS:
        path = SOURCE_DIR / filename
        markdown = path.read_text(encoding="utf-8")
        doc = build_document(title, version, markdown, path)
        output = OUTPUT_DIR / f"{filename.removesuffix('.md')}.docx"
        doc.save(output)
        print(f"built {output.name}")

    build_combined()


def build_combined():
    combined = Document()
    style_document(combined)
    add_header_footer(combined, "课程大作业完整报告")
    add_cover(combined, "课程大作业完整报告", "V1.0")
    add_toc(combined)
    for index, (filename, title, _) in enumerate(DOCUMENTS):
        path = SOURCE_DIR / filename
        markdown = path.read_text(encoding="utf-8")
        p = combined.add_paragraph(style="Heading 1")
        # Word 更新目录后会留下一个字段尾段落。若首个部分再次强制分页，
        # 会产生一张只有页眉页脚的空白页；后续部分仍强制从新页开始。
        if index:
            p.paragraph_format.page_break_before = True
        add_inline(p, f"第{index + 1}部分  {title}", size=18, bold=True)
        add_markdown(combined, markdown, path, normalize_chapter_heading=True)
    props = combined.core_properties
    props.title = f"{PROJECT_NAME} - 课程大作业完整报告"
    props.subject = COURSE_NAME
    props.author = AUTHOR
    props.keywords = "软件工程, 课程大作业, 完整报告, Vue, Flask"
    props.comments = "包含开发计划、需求、设计、测试、总结、个人考核材料与 API 文档"
    output = OUTPUT_DIR / f"{PROJECT_NAME}-课程大作业完整报告.docx"
    combined.save(output)
    print(f"built {output.name}")


if __name__ == "__main__":
    build_all()
